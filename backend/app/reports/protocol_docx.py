from typing import Any

from docx.enum.text import WD_BREAK

from app.reports.docx_common import (
    FORMAT_LABELS,
    LEVEL_LABELS,
    ORGANIZATION,
    PLAN_LABELS,
    STATUS_LABELS,
    add_table,
    fmt_date,
    fmt_num,
    fmt_time,
    footer_with_page_number,
    new_document,
    para,
    signature_block,
    to_bytes,
)


def build_protocol_docx(info: dict[str, Any]) -> bytes:
    """Официальный протокол: титульный лист, таблица результатов (все участники), судейская коллегия, подписи.

    info: competition (title, discipline, type, level, plan, plan_number, start/end, location, format,
    age_groups, stages), protocol (number, version, published_at), rows, judges.
    """
    comp, protocol, rows, judges = info["competition"], info["protocol"], info["rows"], info["judges"]
    doc = new_document()
    footer_with_page_number(doc, f"Протокол № {protocol['number']}")

    # --- титульный лист ---
    para(doc, ORGANIZATION.upper(), bold=True, size=12, align="center", after=60)
    para(doc, f"ПРОТОКОЛ № {protocol['number']}", bold=True, size=20, align="center", after=6)
    para(doc, "результатов соревнования", size=14, align="center", after=18)
    para(doc, comp["title"], bold=True, size=16, align="center", after=18)
    para(doc, f"Вид спорта: спортивное программирование", size=12, align="center", after=2)
    para(doc, f"Дисциплина: {comp['discipline']}", size=12, align="center", after=2)
    para(doc, f"Дата проведения: {fmt_date(comp['start_date'])} – {fmt_date(comp['end_date'])}", size=12, align="center", after=2)
    para(doc, f"Место проведения: {comp['location'] or '–'}", size=12, align="center", after=24)

    plan = PLAN_LABELS[comp["plan"]] + (f", № {comp['plan_number']}" if comp["plan_number"] else "")
    details = [
        ["Тип и статус", f"{comp['type']}, {LEVEL_LABELS[comp['level']]}"],
        ["Включено в календарный план", plan],
        ["Формат проведения", {"online": "онлайн", "offline": "очно", "mixed": "смешанный (онлайн и очно)"}[comp["format"]]],
        ["Возрастные группы", ", ".join(comp["age_groups"])],
        ["Участников в протоколе", str(len(rows))],
        ["Протокол опубликован", f"{fmt_date(protocol['published_at'])}, версия {protocol['version']}"],
    ]
    add_table(doc, ["Параметр", "Значение"], details, [5.5, 11.0], font_size=11)
    if comp["stages"]:
        para(doc, "Этапы соревнования", bold=True, size=12, before=14, after=4)
        add_table(doc, ["№", "Этап", "Формат", "Даты", "Место / ссылка"],
                  [[str(s["order_no"]), s["name"], FORMAT_LABELS[s["format"]],
                    f"{s['starts_at']:%d.%m.%Y %H:%M} – {s['ends_at']:%d.%m.%Y %H:%M}", s["place_or_url"] or "–"] for s in comp["stages"]],
                  [1.0, 5.2, 2.0, 4.8, 3.5], font_size=9, center_cols={0, 2})
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- результаты ---
    para(doc, f"Результаты соревнования", bold=True, size=14, align="center", after=2)
    para(doc, comp["title"], italic=True, size=11, align="center", after=10)
    has_team = any(r.get("team") for r in rows)
    has_tasks = any(r.get("tasks_solved") is not None for r in rows)
    headers = ["Место", "Фамилия, имя, отчество", "Дата рожд."]
    widths = [1.4, 5.6, 2.2]
    if has_team:
        headers.append("Команда")
        widths.append(3.0)
    headers.append("Результат (баллы)")
    widths.append(2.2)
    if has_tasks:
        headers.append("Решено задач")
        widths.append(1.8)
    headers += ["Время", "Примечание"]
    widths += [2.0, 2.6]
    # подгоняем суммарную ширину под страницу
    scale = 17.0 / sum(widths)
    widths = [w * scale for w in widths]

    table_rows = []
    for r in rows:
        line = [str(r["place"]) if r["place"] else "–", r["full_name"], fmt_date(_date(r["birth_date"]))]
        if has_team:
            line.append(r.get("team") or "–")
        line.append(fmt_num(r["score"]) if r["score"] is not None else "–")
        if has_tasks:
            line.append(str(r["tasks_solved"]) if r["tasks_solved"] is not None else "–")
        line += [fmt_time(r["time_seconds"]) or "–", STATUS_LABELS.get(r["status"], "") or (r.get("note") or "")]
        table_rows.append(line)
    add_table(doc, headers, table_rows, widths, font_size=9.5, center_cols={0, 2} | set(range(3 + has_team, len(headers) - 1)))
    para(doc, f"Всего участников: {len(rows)}. В протокол включены все участники с одобренными заявками, "
              f"места присвоены по баллам, при равенстве – по числу решённых задач и времени.", italic=True, size=10, before=6)

    # --- судейская коллегия ---
    para(doc, "Судейская коллегия", bold=True, size=14, align="center", before=16, after=6)
    if judges:
        add_table(doc, ["№", "Судейская должность", "Фамилия, имя, отчество", "Судейская категория"],
                  [[str(i), j["position"], j["name"], j["category"] or "–"] for i, j in enumerate(judges, start=1)],
                  [1.0, 5.0, 6.5, 4.5], font_size=10, center_cols={0})
    else:
        para(doc, "Судейская коллегия на соревнование не назначена.", italic=True)

    # --- подписи ---
    para(doc, "", after=10)
    chief = next((j["name"] for j in judges if j["position"] == "Главный судья"), "")
    secretary = next((j["name"] for j in judges if j["position"] == "Главный секретарь"), "")
    signature_block(doc, "Главный судья", chief)
    signature_block(doc, "Главный секретарь", secretary)
    signature_block(doc, "Председатель регионального отделения")
    para(doc, "М.П.", size=10, before=8)
    para(doc, "Документ сформирован автоматически из данных платформы ФСП Оренбургской области.", italic=True, size=9, before=6)
    return to_bytes(doc)


def _date(value):
    from datetime import date

    return value if isinstance(value, date) else date.fromisoformat(value)
