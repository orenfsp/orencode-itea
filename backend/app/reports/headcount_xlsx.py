from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.headcount import HeadcountReport, Row

NAVY, LIGHT, ZEBRA = "1F3A5F", "D9E2F3", "F5F8FC"
THIN = Side(style="thin", color="9AA7B8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

ORGANIZATION = "Региональное отделение Федерации спортивного программирования Оренбургской области"


def _title(ws: Worksheet, title: str, subtitle: str, width: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A1"], ws["A2"] = title, subtitle
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color="55657A")
    ws["A1"].alignment = ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    return 4


def _header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=text)
        c.font, c.fill, c.border, c.alignment = HEADER_FONT, HEADER_FILL, BORDER, CENTER
    ws.row_dimensions[row].height = 30


def _delta(first: int, last: int) -> tuple[int, float | str]:
    return last - first, (round((last - first) / first * 100, 1) if first else "–")


def _dynamics_table(ws: Worksheet, start: int, first_header: str, years: list[int], rows: list[Row],
                    total: dict[int, int] | None, total_label: str, with_extra: str | None = None) -> int:
    """Таблица «объект × годы» с динамикой и итоговой строкой. Возвращает номер строки после таблицы."""
    headers = [first_header, *([with_extra] if with_extra else []), *map(str, years), f"Δ {years[0]}→{years[-1]}, чел.", "Δ, %"]
    _header(ws, start, headers)
    r = start + 1
    for i, row in enumerate(rows):
        values = [row.label, *([row.extra] if with_extra else []), *(row.counts[y] for y in years)]
        d, pct = _delta(row.counts[years[0]], row.counts[years[-1]])
        for col, v in enumerate([*values, d, pct], start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = BORDER
            c.alignment = LEFT if col <= (2 if with_extra else 1) else CENTER
            if i % 2:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
            if col > (2 if with_extra else 1) + len(years):
                c.number_format = '+0;-0;0' if col == len(headers) - 1 else '+0.0"%";-0.0"%";0.0"%"'
        r += 1
    if total is not None:
        values = [total_label, *([""] if with_extra else []), *(total[y] for y in years)]
        d, pct = _delta(total[years[0]], total[years[-1]])
        for col, v in enumerate([*values, d, pct], start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font, c.fill, c.border = Font(bold=True), TOTAL_FILL, BORDER
            c.alignment = LEFT if col <= (2 if with_extra else 1) else CENTER
            if col > (2 if with_extra else 1) + len(years):
                c.number_format = '+0;-0;0' if col == len(headers) - 1 else '+0.0"%";-0.0"%";0.0"%"'
        r += 1
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    return r + 1


def _note(ws: Worksheet, row: int, text: str, width: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.alignment = Font(italic=True, size=9, color="55657A"), Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 15 * (1 + len(text) // 110)
    return row + 1


def _autowidth(ws: Worksheet, minimum: int = 10, maximum: int = 58) -> None:
    merged = {cell for rng in ws.merged_cells.ranges for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row,
                                                                              min_col=rng.min_col, max_col=rng.max_col) for cell in row}
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None or c in merged:
                continue
            widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(maximum, int(w * 1.15) + 2))


def _page(ws: Worksheet) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False


def build_headcount_xlsx(report: HeadcountReport, generated_at: datetime) -> bytes:
    years = report.years
    period = f"{years[0]}–{years[-1]}"
    stamp = f"Сформировано {generated_at:%d.%m.%Y %H:%M}. {ORGANIZATION}"
    wb = Workbook()

    # --- Сводка ---
    ws = wb.active
    ws.title = "Сводка"
    width = len(years) + 3
    r = _title(ws, f"Численность занимающихся спортивным программированием в Оренбургской области, {period}", stamp, width)
    summary = [
        Row("Занимающихся спортсменов (уникальных)", report.total),
        Row("из них впервые начали заниматься в году", report.new),
        Row("Тренеры, ведущие подопечных", report.coaches),
        Row("Судьи с действующей категорией", report.judges),
    ]
    end = _dynamics_table(ws, r, "Показатель", years, summary, None, "")
    end = _note(ws, end, "Спортсмен считается занимающимся в году, если его занятия по дисциплине пересекаются с календарным годом. "
                         "Данные формируются из реестра занимающихся платформы (занятия, организации, дисциплины).", width)
    chart = BarChart()
    chart.type, chart.title, chart.style = "col", "Численность занимающихся по годам", 10
    chart.y_axis.title, chart.x_axis.title = "человек", "год"
    first_col = 2
    data = Reference(ws, min_col=first_col, max_col=first_col + len(years) - 1, min_row=r + 1, max_row=r + 1)
    chart.add_data(data, from_rows=True, titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=first_col, max_col=first_col + len(years) - 1, min_row=r, max_row=r))
    chart.legend = None
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, f"A{end + 1}")
    _autowidth(ws)
    ws.column_dimensions["A"].width = 48
    _page(ws)

    # --- Муниципалитеты ---
    ws = wb.create_sheet("Муниципалитеты")
    width = len(years) + 3
    r = _title(ws, f"Занимающиеся по муниципалитетам, {period}", stamp, width)
    end = _dynamics_table(ws, r, "Муниципальное образование", years, report.municipalities, report.total, "Итого по области")
    _note(ws, end, "Муниципалитет определяется по месту проживания спортсмена; каждый спортсмен учитывается один раз.", width)
    _autowidth(ws)
    _page(ws)

    # --- Организации ---
    ws = wb.create_sheet("Организации")
    width = len(years) + 4
    r = _title(ws, f"Занимающиеся по организациям (вузы, колледжи, школы, спортшколы, клубы), {period}", stamp, width)
    end = _dynamics_table(ws, r, "Организация", years, report.organizations, report.total, "Итого (уникальных спортсменов)", with_extra="Тип организации")
    _note(ws, end, "Спортсмен, занимающийся в нескольких организациях, учитывается в каждой, поэтому сумма строк может превышать итог.", width)
    _autowidth(ws)
    _page(ws)

    # --- Дисциплины ---
    ws = wb.create_sheet("Дисциплины")
    width = len(years) + 3
    r = _title(ws, f"Занимающиеся по дисциплинам вида спорта, {period}", stamp, width)
    end = _dynamics_table(ws, r, "Дисциплина", years, report.disciplines, report.total, "Итого (уникальных спортсменов)")
    _note(ws, end, "Спортсмен, занимающийся несколькими дисциплинами, учитывается в каждой, поэтому сумма строк может превышать итог.", width)
    _autowidth(ws)
    _page(ws)

    # --- Возрастные группы ---
    ws = wb.create_sheet("Возрастные группы")
    width = len(years) + 3
    r = _title(ws, f"Занимающиеся по возрастным группам, {period}", stamp, width)
    end = _dynamics_table(ws, r, "Возрастная группа", years, report.age_groups, report.total, "Итого")
    _note(ws, end, "Возраст считается на 31 декабря соответствующего года.", width)
    _autowidth(ws)
    _page(ws)

    # --- Дисциплины × возраст ---
    ws = wb.create_sheet("Дисциплины × возраст")
    cols = report.matrix_columns
    width = len(cols) + 2
    r = _title(ws, f"Дисциплины по возрастным группам, {report.matrix_year} год", stamp, width)
    _header(ws, r, ["Дисциплина", *cols, "Итого"])
    for i, (name, counts) in enumerate(report.matrix):
        row = r + 1 + i
        for col, v in enumerate([name, *(counts[c] for c in cols), sum(counts.values())], start=1):
            c = ws.cell(row=row, column=col, value=v)
            c.border, c.alignment = BORDER, LEFT if col == 1 else CENTER
            if i % 2:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
    row = r + 1 + len(report.matrix)
    totals = [sum(counts[c] for _, counts in report.matrix) for c in cols]
    for col, v in enumerate(["Итого (спортсмено-дисциплин)", *totals, sum(totals)], start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font, c.fill, c.border, c.alignment = Font(bold=True), TOTAL_FILL, BORDER, LEFT if col == 1 else CENTER
    ws.freeze_panes = ws.cell(row=r + 1, column=2)
    _autowidth(ws)
    _page(ws)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
