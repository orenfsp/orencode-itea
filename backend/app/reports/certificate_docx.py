from io import BytesIO
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from app.reports.docx_common import (
    LEVEL_LABELS,
    ORGANIZATION,
    footer_with_page_number,
    fmt_date,
    fmt_num,
    new_document,
    para,
    signature_block,
    to_bytes,
)


def build_certificate_docx(info: dict[str, Any], qr: bytes) -> bytes:
    """Справка об участии в соревновании с номером и QR-кодом проверки подлинности."""
    doc = new_document()
    footer_with_page_number(doc, f"Справка № {info['number']}")

    para(doc, ORGANIZATION.upper(), bold=True, size=12, align="center", after=18)
    para(doc, "СПРАВКА", bold=True, size=20, align="center", after=2)
    para(doc, "об участии в соревновании", size=14, align="center", after=4)
    para(doc, f"№ {info['number']}", bold=True, size=13, align="center", after=18)

    para(doc, f"Настоящая справка выдана {info['athlete_name']}, {fmt_date(info['birth_date'])} года рождения, "
              f"в том, что он(а) принял(а) участие в соревновании по спортивному программированию:", align="justify", after=8)
    para(doc, f"«{info['competition_title']}»", bold=True, size=13, align="center", after=8)
    rows = [
        ("Дисциплина", info["discipline"]),
        ("Статус соревнования", LEVEL_LABELS[info["level"]]),
        ("Даты проведения", f"{fmt_date(info['start_date'])} – {fmt_date(info['end_date'])}"),
        ("Место проведения", info["location"] or "–"),
        ("Результат", f"{info['place']} место из {info['participants']} участников"),
    ]
    if info.get("score") is not None:
        rows.append(("Набрано баллов", fmt_num(info["score"])))
    rows.append(("Протокол", info["protocol_number"]))
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, str(value)
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].width, cells[1].width = Cm(5.5), Cm(11)

    para(doc, "", after=10)
    para(doc, "Справка выдана для предъявления по месту требования (портфолио, поступление, стипендия).", align="justify")
    para(doc, f"Дата выдачи: {fmt_date(info['issued_at'])}", after=14)

    # блок проверки подлинности: QR слева, пояснение справа
    box = doc.add_table(rows=1, cols=2)
    left, right = box.rows[0].cells
    left.width, right.width = Cm(4.2), Cm(12.3)
    left.paragraphs[0].add_run().add_picture(BytesIO(qr), width=Cm(3.6))
    right.text = ""
    p = right.paragraphs[0]
    p.add_run("Проверка подлинности\n").bold = True
    p.add_run("Отсканируйте QR-код или откройте ссылку: подлинная справка отображается на публичной странице платформы "
              "с ФИО, соревнованием, дисциплиной и занятым местом.\n")
    p.add_run(info["verify_url"]).font.size = None
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    signature_block(doc, "Председатель регионального отделения")
    para(doc, "М.П.", size=10, before=6)
    return to_bytes(doc)
